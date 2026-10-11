"""Build small DOCX / PPTX / XLSX / MD files for citation-highlight checks.

The files are generated, never committed. Run from ``backend/``::

    python3 scripts/make_highlight_fixtures.py OUT_DIR          # write the files
    python3 scripts/make_highlight_fixtures.py OUT_DIR --json   # also dump the
        # pages, chunks and citation payloads the app would produce for them

The JSON dump runs the real extractors and ``ParseService.chunk_document``.
``pages[].text`` is what ``pages.content`` stores. Each citation carries
``text_snippet`` and ``context_text`` built with the formulas of
``chat_service._citation_payload``, the real ``focus_sentence`` for a claim
equal to the chunk's longest sentence, and the page range from
``citation_location``. ``tests/test_extractors_office.py`` imports the builders.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import sys
from pathlib import Path
from typing import Any

# Make the backend root importable when run as `python3 scripts/make_highlight_fixtures.py`.
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

REPEATED_SENTENCE = "All analyses followed the intention-to-treat principle."

DOCX_SECTIONS: list[tuple[str, list[str]]] = [
    ("Study design", [
        "The trial enrolled 412 adults across six regional clinics between March 2021 and June 2022.",
        "Participants were randomised to the intervention or to usual care in a one-to-one ratio.",
        REPEATED_SENTENCE,
    ]),
    ("Outcomes", [
        "The primary outcome was the change in systolic blood pressure at twelve months.",
        "Secondary outcomes included adherence, adverse events and quality-of-life scores.",
        REPEATED_SENTENCE,
    ]),
]
DOCX_TABLE = [
    ["Site", "Enrolled", "Completed"],
    ["Northfield", "88", "81"],
    ["Eastbrook", "74", "70"],
]

PPTX_SLIDES: list[dict[str, Any]] = [
    {
        "title": "Quarterly review",
        "body": "Revenue grew 14 percent year over year, driven by the enterprise segment.",
        "notes": "Mention the two delayed contracts before the questions.",
    },
    {
        "title": "Regional results",
        "table": [
            ["Region", "Revenue", "Growth"],
            ["North", "4.2M", "12%"],
            ["South", "3.1M", "9%"],
        ],
    },
    {
        "title": "Quarterly review",
        "body": "Operating costs fell for the third consecutive quarter.",
    },
]

XLSX_SHEETS: dict[str, list[list[Any]]] = {
    "Budget": [
        ["Item", "Q1", "Q2"],
        ["Travel", 1200, 950],
        ["Software", 4300, 4300],
    ],
    "Headcount": [
        ["Team", "People"],
        ["Research", 12],
        [None, None],
        ["Support", 7],
    ],
}

MD_TEXT = """# Field notes

The survey covered 38 wetlands in the northern basin during the 2023 season.
Water samples were taken at dawn to limit the effect of daily temperature swings.

| Site | Depth (m) | pH |
| --- | --- | --- |
| Alder Marsh | 1.4 | 6.8 |
| Reed Hollow | 0.9 | 7.1 |

Use `--strict` mode \\*only\\* when the probe is calibrated &amp; dry.

## 中文摘要

本研究共调查了三十八处湿地。水样均在清晨采集，以减少昼夜温差的影响。

## ملخص

شملت الدراسة ثمانية وثلاثين موقعا من الأراضي الرطبة. جمعت العينات عند الفجر.
"""


def build_docx() -> bytes:
    """Two Heading 1 sections; the first holds a 3x3 table."""
    from docx import Document

    doc = Document()
    for index, (heading, paragraphs) in enumerate(DOCX_SECTIONS):
        doc.add_heading(heading, level=1)
        for text in paragraphs:
            doc.add_paragraph(text)
        if index == 0:
            table = doc.add_table(rows=len(DOCX_TABLE), cols=len(DOCX_TABLE[0]))
            for r, row in enumerate(DOCX_TABLE):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pptx() -> bytes:
    """Three slides: title + body + notes, title + table, a repeated title."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for spec in PPTX_SLIDES:
        if "table" in spec:
            slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
            slide.shapes.title.text = spec["title"]
            rows = spec["table"]
            shape = slide.shapes.add_table(
                len(rows), len(rows[0]), Inches(1), Inches(2), Inches(8), Inches(2)
            )
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = value
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
            slide.shapes.title.text = spec["title"]
            slide.placeholders[1].text_frame.text = spec["body"]
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_xlsx() -> bytes:
    """Two sheets; the second has an empty row that the extractor skips."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in XLSX_SHEETS.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_md() -> bytes:
    """Headings, a GFM table, escapes/entities/inline code, zh and ar sections."""
    return MD_TEXT.encode("utf-8")


BUILDERS = {
    "docx": build_docx,
    "pptx": build_pptx,
    "xlsx": build_xlsx,
    "md": build_md,
}


def page_infos(extracted: list) -> list:
    """Mirror parse_worker's non-PDF ExtractedPage -> PageInfo conversion."""
    from app.services.parse_service import BlockInfo, PageInfo

    pages = []
    for ep in extracted:
        blocks = []
        if ep.section_title:
            blocks.append(BlockInfo(
                page=ep.page_number,
                text=ep.section_title,
                bbox=(0.0, 0.0, 1.0, 0.08),
                font_size=18.0,
            ))
        blocks.append(BlockInfo(
            page=ep.page_number,
            text=ep.text,
            bbox=(0.0, 0.1 if ep.section_title else 0.0, 1.0, 1.0),
            font_size=12.0,
        ))
        pages.append(PageInfo(
            page_number=ep.page_number,
            width_pt=ep.width_pt or 612.0,
            height_pt=ep.height_pt or 792.0,
            rotation=0,
            blocks=blocks,
        ))
    return pages


def citation_like(chunk) -> dict[str, Any]:
    """The citation fields the reader receives for ``chunk``."""
    from app.services.citation_focus_service import _split_sentences, focus_sentence
    from app.services.citation_location import citation_location

    text = chunk.text or ""
    sentences = _split_sentences(text)
    claim = max(sentences, key=len) if sentences else ""
    focus = focus_sentence(text, claim) if claim else None
    page_end = chunk.page_end if (chunk.page_end and chunk.page_end != chunk.page_start) else None
    prefix = f"[page {chunk.page_start}{'–' + str(page_end) if page_end else ''}] " if chunk.page_start else ""
    payload: dict[str, Any] = {
        "text_snippet": ((f"{chunk.section_title}: " if chunk.section_title else "") + text)[:100],
        "context_text": (prefix + text)[:900],
        "claim": claim,
    }
    if focus:
        payload["focus_snippet"] = focus
    location = citation_location(
        chunk.page_start, chunk.page_end, chunk.bboxes, source_text=text, focus=focus or ""
    )
    payload["page"] = location["page"]
    payload["page_end"] = location["page_end"]
    return payload


def pipeline(file_type: str, data: bytes) -> dict[str, Any]:
    """Extract, chunk and build citation payloads like the parse worker and chat."""
    from app.services.extractors import extract_document
    from app.services.parse_service import ParseService

    extracted = extract_document(data, file_type)
    chunks = ParseService().chunk_document(page_infos(extracted))
    return {
        "file_type": file_type,
        "pages": [
            {"page_number": ep.page_number, "section_title": ep.section_title, "text": ep.text}
            for ep in extracted
        ],
        "chunks": [
            {
                "chunk_index": c.chunk_index,
                "page_start": c.page_start,
                "page_end": c.page_end,
                "section_title": c.section_title,
                "text": c.text,
                "bboxes": c.bboxes,
                "citation": citation_like(c),
            }
            for c in chunks
        ],
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("out_dir", type=Path)
    parser.add_argument("--json", action="store_true", help="also write <type>.pipeline.json")
    args = parser.parse_args(argv)
    args.out_dir.mkdir(parents=True, exist_ok=True)
    for file_type, build in BUILDERS.items():
        data = build()
        target = args.out_dir / f"highlight-fixture.{file_type}"
        target.write_bytes(data)
        print(f"wrote {target} ({len(data)} bytes)")
        if args.json:
            dump = args.out_dir / f"highlight-fixture.{file_type}.pipeline.json"
            dump.write_text(json.dumps(pipeline(file_type, data), ensure_ascii=False, indent=2))
            print(f"wrote {dump}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
