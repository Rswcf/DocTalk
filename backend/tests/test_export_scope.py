"""An exported beyond-document answer carries the same label as in the app and on a shared page: an export is
passed on, and a general-knowledge answer must never read as one drawn from the document."""
from __future__ import annotations

from types import SimpleNamespace

from app.services.export_service import render_docx, render_markdown

LABEL = "Answered from general knowledge — not verified against the document"


def _conversation(answer_meta):
    return [
        SimpleNamespace(role="user", content="Who wrote this novel?", citations=None, metadata_json=None),
        SimpleNamespace(role="assistant", content="A nineteenth-century writer.", citations=None,
                        metadata_json=answer_meta),
    ]


def test_markdown_export_labels_a_beyond_answer():
    md = render_markdown("Novel", "novel.pdf", _conversation({"answer_scope": "beyond_document"}))

    assert f"**A:** *{LABEL}*\n\nA nineteenth-century writer." in md


def test_markdown_export_of_a_grounded_answer_is_unchanged():
    md = render_markdown("Novel", "novel.pdf", _conversation({}))

    assert "**A:** A nineteenth-century writer." in md
    assert "general knowledge" not in md


def test_docx_export_labels_a_beyond_answer_and_only_that():
    from docx import Document

    beyond = [p.text for p in Document(render_docx("Novel", "novel.pdf", _conversation({"answer_scope": "beyond_document"}))).paragraphs]
    grounded = [p.text for p in Document(render_docx("Novel", "novel.pdf", _conversation({}))).paragraphs]

    assert beyond.index(LABEL) == beyond.index("A nineteenth-century writer.") - 1
    assert LABEL not in grounded
