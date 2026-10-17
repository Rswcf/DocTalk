"""FE-04: real PostgreSQL persistence through HTTP reads, shares and exports.

Run only in the isolated integration-test database provisioned by conftest.
No LLM, object storage, production service or browser is involved. Only auth
and anonymous rate limiting are stubbed; SQL, routes and renderers are real.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock

import fitz
import pytest
import pytest_asyncio
from docx import Document as DocxDocument
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy import delete

from app.api import chat, export, sharing
from app.core.deps import get_current_user_optional, get_db_session, require_auth
from app.models.database import AsyncSessionLocal
from app.models.tables import (
    AnswerShare,
    ChatSession,
    Chunk,
    Document,
    Message,
    Page,
    User,
)
from app.services.chat_service import (
    _ChunkInfo,
    _citation_payload,
    _hydrate_citation_pages,
)

pytestmark = [pytest.mark.integration, pytest.mark.asyncio(loop_scope="session")]
FIRST = "Revenue for 2026 was EUR 1,200,000. Operating costs were EUR 800,000. "
MIDDLE = "The board approved an irrigation project at its annual meeting."
LAST = "This document does not state the company founding year."


def digest(snapshot):
    return hashlib.sha256(json.dumps(snapshot, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


@pytest_asyncio.fixture(loop_scope="session")
async def evidence_case(monkeypatch):
    monkeypatch.setattr(sharing.shared_view_limiter, "is_allowed", AsyncMock(return_value=True))
    async with AsyncSessionLocal() as db:
        owner = User(email=f"citation-projection-{uuid.uuid4().hex}@example.com", plan="pro")
        db.add(owner)
        await db.flush()
        doc = Document(filename="evidence.pdf", file_size=1, storage_key="qa-citation-projection",
                       user_id=owner.id, status="ready")
        db.add(doc)
        await db.flush()
        texts = {1: FIRST * 16, 2: MIDDLE, 3: LAST}
        db.add_all([Page(document_id=doc.id, page_number=n, content=text) for n, text in texts.items()])
        chunk = Chunk(document_id=doc.id, chunk_index=0, text=" ".join(texts.values()),
                      token_count=400, page_start=1, page_end=3, bboxes=[], section_title="Financial results")
        session = ChatSession(document_id=doc.id, user_id=owner.id, title="Founding year evidence")
        db.add_all([chunk, session])
        await db.flush()
        source = _ChunkInfo(id=chunk.id, document_id=doc.id, document_filename=doc.filename,
                            page_start=1, page_end=3, bboxes=[], text=chunk.text,
                            section_title=chunk.section_title)
        await _hydrate_citation_pages(db, {1: source})
        citation = _citation_payload(1, source, len("The founding year is not stated."), LAST)
        assert citation["text_snippet"] == LAST
        assert citation["page"] == citation["page_end"] == 3
        assert LAST not in citation["context_text"]  # source context is capped at 900
        # Persist the exact historical defect: correct verified focus and page,
        # wrong prefix excerpt. Projection must repair every subsequent read.
        citation["text_snippet"] = "Financial results: " + FIRST[:80]
        message = Message(session_id=session.id, role="assistant",
                          content="The founding year is not stated.", citations=[citation])
        db.add(message)
        await db.commit()
        case = SimpleNamespace(owner=owner.id, document=doc.id, session=session.id,
                               message=message.id, original=copy.deepcopy(citation))
    app = FastAPI()
    app.include_router(chat.chat_router)
    app.include_router(sharing.router)
    app.include_router(export.router)

    async def database():
        async with AsyncSessionLocal() as db:
            yield db

    async def actor():
        async with AsyncSessionLocal() as db:
            return await db.get(User, case.owner)

    app.dependency_overrides[get_db_session] = database
    app.dependency_overrides[get_current_user_optional] = actor
    app.dependency_overrides[require_auth] = actor
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            case.client = client
            yield case
    finally:
        async with AsyncSessionLocal() as db:
            await db.execute(delete(Document).where(Document.id == case.document))
            await db.execute(delete(User).where(User.id == case.owner))
            await db.commit()


async def assert_stored_citation_unchanged(case):
    async with AsyncSessionLocal() as db:
        record = await db.get(Message, case.message)
        assert record.citations == [case.original]


async def test_history_read_projects_verified_focus_without_rewriting_jsonb(evidence_case):
    case = evidence_case
    response = await case.client.get(f"/api/sessions/{case.session}/messages")
    assert response.status_code == 200, response.text
    citation = response.json()["messages"][0]["citations"][0]
    assert citation["text_snippet"] == citation["focus_snippet"] == LAST
    assert citation["page"] == citation["page_end"] == 3
    await assert_stored_citation_unchanged(case)


@pytest.mark.parametrize("format", ["md", "docx", "pdf"])
async def test_http_exports_render_verified_last_page_excerpt(evidence_case, format):
    case = evidence_case
    response = await case.client.get(f"/api/sessions/{case.session}/export", params={"format": format})
    assert response.status_code == 200, response.text[:200]
    assert "attachment" in response.headers["content-disposition"]
    if format == "md":
        text = response.text
    elif format == "docx":
        document = DocxDocument(io.BytesIO(response.content))
        text = "\n".join(p.text for p in document.paragraphs)
    else:
        with fitz.open(stream=response.content, filetype="pdf") as pdf:
            text = "\n".join(page.get_text() for page in pdf)
    assert LAST in text
    assert "Revenue for 2026" not in text
    assert "Financial results" not in text
    assert "Page 3" in text
    await assert_stored_citation_unchanged(case)


async def test_share_preview_digest_create_public_and_existing_immutable_snapshot(evidence_case):
    case = evidence_case
    path = f"/api/sessions/{case.session}/answers/{case.message}/share"
    response = await case.client.get(path)
    assert response.status_code == 200, response.text
    state = response.json()
    preview = state["preview"]
    citation = preview["messages"][0]["citations"][0]
    assert citation["text_snippet"] == LAST
    assert citation["page"] == citation["page_end"] == 3
    assert state["snapshot_digest"] == digest(preview)

    # Model a link published before FE-04; reading it must retain its original
    # wrong excerpt. A new preview/create must produce a distinct corrected link.
    old_snapshot = copy.deepcopy(preview)
    old_snapshot["messages"][0]["citations"][0]["text_snippet"] = case.original["text_snippet"]
    async with AsyncSessionLocal() as db:
        old = AnswerShare(user_id=case.owner, session_id=case.session, message_id=case.message,
                          snapshot=old_snapshot, snapshot_digest=digest(old_snapshot))
        db.add(old)
        await db.commit()
        old_token = str(old.share_token)
    created = await case.client.post(path, json={"snapshot_digest": state["snapshot_digest"]})
    assert created.status_code == 200, created.text
    token = created.json()["share_token"]
    assert token != old_token
    public = await case.client.get(f"/api/shared/{token}")
    assert public.status_code == 200
    assert public.json() == preview
    for private in ("focus_snippet", "context_text", "chunk_id", "bboxes"):
        assert private not in public.text
    old_public = await case.client.get(f"/api/shared/{old_token}")
    assert old_public.status_code == 200
    assert old_public.json() == old_snapshot
    await assert_stored_citation_unchanged(case)

    # Changing the supporting evidence changes the visible preview and digest.
    # Previously consented snapshots remain immutable; stale consent is rejected.
    async with AsyncSessionLocal() as db:
        message = await db.get(Message, case.message)
        updated = copy.deepcopy(message.citations)
        updated[0].update(focus_snippet=MIDDLE, page=2, page_end=2)
        message.citations = updated
        await db.commit()
    changed = await case.client.get(path)
    assert changed.json()["preview"]["messages"][0]["citations"][0]["text_snippet"] == MIDDLE
    assert changed.json()["snapshot_digest"] != state["snapshot_digest"]
    stale = await case.client.post(path, json={"snapshot_digest": state["snapshot_digest"]})
    assert stale.status_code == 409
    assert stale.json()["detail"]["error"] == "SHARE_CHANGED"
    assert (await case.client.get(f"/api/shared/{token}")).json() == preview
    assert (await case.client.get(f"/api/shared/{old_token}")).json() == old_snapshot
