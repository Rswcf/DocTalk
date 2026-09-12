"""Use the same offset/source cases as the browser Markdown regression suite."""
import copy
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from app.services.export_service import (
    _prepare_export,
    render_docx,
    render_markdown,
    render_pdf,
)

CASES = json.loads((Path(__file__).parent / "fixtures/export_citations.json").read_text())


@pytest.mark.parametrize("case", CASES, ids=lambda case: case["name"])
def test_all_sources_and_answer_offsets_survive_markdown_and_docx(case):
    original = copy.deepcopy(case["messages"])
    messages = [SimpleNamespace(**message) for message in case["messages"]]
    prepared, refs = _prepare_export(messages, markdown=True)
    assert [text for _, text in prepared] == case["answers"]
    assert [ref.get("document_filename") for ref in refs] == case["sources"]
    md = render_markdown("Citation QA", "fixture", messages)
    for answer in case["answers"]:
        assert answer in md
    assert re.findall(r"^\[\^(\d+)\]:", md, re.M) == [str(n + 1) for n in range(len(refs))]
    doc = Document(render_docx("Citation QA", "fixture", messages))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    for answer in case["answers"]:
        assert any(answer.replace("[^", "[") in paragraph for paragraph in paragraphs)
    for i, source in enumerate(case["sources"], 1):
        assert any(paragraph.startswith(f"[{i}] ") and source in paragraph for paragraph in paragraphs)
    assert case["messages"] == original


def test_pdf_contains_linked_answer_numbers_and_both_sources():
    import fitz

    messages = [SimpleNamespace(**message) for message in CASES[0]["messages"]]
    pdf = fitz.open(stream=render_pdf("Citation QA", "fixture", messages).getvalue(), filetype="pdf")
    text = "\n".join(page.get_text() for page in pdf)
    assert "A 42.[1]" in text and "B 84.[2]" in text
    assert '[1] Page 1, alpha.pdf, "42"' in text
    assert '[2] Page 2, beta.pdf, "84"' in text


def test_xml_sanitization_happens_after_offset_restoration():
    messages = [SimpleNamespace(role="assistant", content="A\x00B.", citations=[
        {"ref_index": 1, "offset": 4, "page": 1, "text_snippet": "AB"},
    ])]
    doc = Document(render_docx("QA", "fixture", messages))
    assert "AB.[1]" in [paragraph.text for paragraph in doc.paragraphs]
