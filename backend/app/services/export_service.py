"""Render chat sessions as Markdown, DOCX, or PDF."""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any, List

from markupsafe import escape as html_escape

_MAX_MESSAGES = 500

# Codepoints forbidden by XML 1.0 Char production (https://www.w3.org/TR/xml/#charsets):
#   C0 controls (except \t\n\r): \x00-\x08, \x0B, \x0C, \x0E-\x1F
#   Unpaired surrogates:         \uD800-\uDFFF
#   Non-characters:              \uFFFE, \uFFFF
# python-docx rejects these with ValueError; weasyprint+lxml can misbehave similarly.
_XML_INVALID_RE = re.compile(
    r"[\x00-\x08\x0B\x0C\x0E-\x1F\uD800-\uDFFF\uFFFE\uFFFF]"
)


def _sanitize_xml_text(s: Any) -> str:
    """Coerce to str and strip XML-incompatible codepoints.

    None/non-str inputs fall back to '' / str(s) as a last-line defense —
    upstream callers (export.py) are expected to default None already.
    """
    if s is None:
        return ""
    return _XML_INVALID_RE.sub("", str(s))


_LEGACY_CITATION_RE = re.compile(
    r"!?\[[^\]\n]*\]\([^\n)]*\)"
    r"|(?<![\\!])\[(\d+)\](?!\()"
)


def _code_ranges(text: str) -> list[tuple[int, int]]:
    """Conservatively protect indented code and variable-length Markdown fences."""
    ranges: list[tuple[int, int]] = []
    fence: str | None = None
    offset = 0
    for line in text.splitlines(keepends=True):
        content = re.sub(r"^(?: {0,3}>[ \t]?)+", "", line)
        content = re.sub(r"^ {0,3}(?:[-+*]|\d+[.)])[ \t]+", "", content)
        run = re.match(r"^ {0,3}(`{3,}|~{3,})", content)
        if fence:
            ranges.append((offset, offset + len(line)))
            if run and run[1][0] == fence[0] and len(run[1]) >= len(fence) and not content[run.end():].strip():
                fence = None
        elif run:
            fence = run[1]
            ranges.append((offset, offset + len(line)))
        elif content.startswith(("    ", "\t")):
            ranges.append((offset, offset + len(line)))
        offset += len(line)
    cursor = 0
    for run in re.finditer(r"`+", text):
        if run.start() < cursor or any(start <= run.start() < end for start, end in ranges):
            continue
        close = re.search(r"(?<!`)" + re.escape(run[0]) + r"(?!`)", text[run.end():])
        if close:
            cursor = run.end() + close.end()
            ranges.append((run.start(), cursor))
    return ranges


def _prepare_export(messages: List[Any], *, markdown: bool = False) -> tuple[list[tuple[str, str]], list[dict]]:
    """Restore server codepoint offsets and give each message/source a global ID.

    Ref indexes restart in every answer. An index alone is never a source key,
    including after continuation reuses an index for a different document.
    Stored content remains untouched; XML sanitization happens after insertion.
    """
    prepared: list[tuple[str, str]] = []
    references: list[dict] = []
    for msg in messages:
        text = msg.content or ""
        raw = msg.citations
        if msg.role != "assistant" or not isinstance(raw, list):
            prepared.append((msg.role, text))
            continue
        citations = [c for c in raw if isinstance(c, dict)
                     and type(c.get("ref_index")) is int and c["ref_index"] > 0]
        source_ids: dict[tuple, int] = {}

        def marker(c):
            key = tuple(str(c.get(field) or "") for field in (
                "ref_index", "document_id", "document_filename", "chunk_id",
                "page", "page_end", "text_snippet",
            ))
            if key not in source_ids:
                source_ids[key] = len(references) + 1
                references.append({**c, "ref_index": source_ids[key]})
            number = source_ids[key]
            return f"[^{number}]" if markdown else f"[{number}]"

        # Apply all edits against the original text, never against a string
        # already lengthened by earlier markers. Same-offset citations retain order.
        edits: list[tuple[int, int, dict]] = []
        legacy: dict[int, dict] = {}
        ambiguous: set[int] = set()
        for c in citations:
            offset = c.get("offset")
            if type(offset) is int and 0 <= offset <= len(text):
                edits.append((offset, offset, c))
            elif "offset" not in c:
                ref = c["ref_index"]
                if ref in legacy and legacy[ref] != c:
                    ambiguous.add(ref)
                legacy[ref] = c
        # Only legacy records without offsets use textual markers. Code, links,
        # escaped brackets and ordinary bracketed numbers remain literal text.
        code_ranges = _code_ranges(text) if legacy else []
        # A numeric Markdown link definition makes [n] a link reference, not
        # reliable evidence of a citation. Preserve every use of that label.
        link_labels = {match[1].strip().lower() for match in re.finditer(r"(?m)^ {0,3}\[([^\]\n]+)\]:", text)}
        for match in re.finditer(r"!?\[([^\]\n]*)\]\s*\[([^\]\n]*)\]", text):
            if (match[2] or match[1]).strip().lower() in link_labels:
                code_ranges.append((match.start(), match.end()))
        for match in _LEGACY_CITATION_RE.finditer(text) if legacy else []:
            if any(start <= match.start() < end for start, end in code_ranges):
                continue
            ref = int(match[1]) if match[1] else None
            if ref in legacy and ref not in ambiguous and str(ref) not in link_labels:
                edits.append((match.start(), match.end(), legacy[ref]))
        pieces: list[str] = []
        cursor = 0
        seen: set[tuple[int, str]] = set()
        for start, end, c in sorted(edits, key=lambda edit: edit[0]):
            if start < cursor:
                continue
            label = marker(c)
            if (start, label) in seen:
                continue
            seen.add((start, label))
            pieces.extend((text[cursor:start], label))
            cursor = end
        pieces.append(text[cursor:])
        prepared.append((msg.role, "".join(pieces)))
    return prepared, references


def _format_footnote(c: dict) -> str:
    page = c.get("page", "?")
    end = c.get("page_end")
    location = f"Pages {page}–{end}" if isinstance(page, int) and isinstance(end, int) and end > page else f"Page {page}"
    doc = c.get("document_filename") or ""
    snippet = str(c.get("text_snippet") or "")
    snippet = snippet[:80] + ("..." if len(snippet) > 80 else "")
    source = f"{doc}, " if doc else ""
    return f"{location}, {source}\"{snippet}\""


def render_markdown(title: str, doc_name: str, messages: List[Any]) -> str:
    if len(messages) > _MAX_MESSAGES:
        raise ValueError(f"Export limited to {_MAX_MESSAGES} messages")

    lines = [
        f"# {title}",
        f"*Document: {doc_name}*",
        f"*Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}*",
        "",
        "---",
        "",
    ]

    prepared, citations = _prepare_export(messages, markdown=True)
    for role, text in prepared:
        if role == "user":
            lines.append(f"**Q:** {text}")
        else:
            lines.append(f"**A:** {text}")
        lines.append("")

    if citations:
        lines.append("---")
        lines.append("## References")
        lines.append("")
        for c in citations:
            idx = c.get("ref_index", 0)
            lines.append(f"[^{idx}]: {_format_footnote(c)}")

    lines.extend(["", "---", "*Generated by DocTalk — www.doctalk.site*"])
    return "\n".join(lines)


class ExportError(Exception):
    """Raised when export rendering fails."""


def render_docx(title: str, doc_name: str, messages: List[Any]) -> io.BytesIO:
    if len(messages) > _MAX_MESSAGES:
        raise ValueError(f"Export limited to {_MAX_MESSAGES} messages")

    try:
        from docx import Document
    except ImportError:
        raise ExportError("DOCX export is not available: python-docx not installed")
    from docx.shared import Pt, RGBColor

    safe_title = _sanitize_xml_text(title)
    safe_doc_name = _sanitize_xml_text(doc_name)

    doc = Document()
    doc.add_heading(safe_title, level=1)
    p = doc.add_paragraph()
    run = p.add_run(
        f"Document: {safe_doc_name} | Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(128, 128, 128)

    prepared, citations = _prepare_export(messages)
    for role, text in prepared:
        content = _sanitize_xml_text(text)
        if role == "user":
            p = doc.add_paragraph()
            run = p.add_run(f"Q: {content}")
            run.bold = True
        else:
            doc.add_paragraph(content)

    if citations:
        doc.add_heading("References", level=2)
        for c in citations:
            idx = c.get("ref_index", 0)
            doc.add_paragraph(f"[{idx}] {_sanitize_xml_text(_format_footnote(c))}")

    p = doc.add_paragraph()
    run = p.add_run("Generated by DocTalk — www.doctalk.site")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(160, 160, 160)

    try:
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf
    except Exception as e:
        raise ExportError(f"DOCX rendering failed: {e}") from e


def render_pdf(title: str, doc_name: str, messages: List[Any]) -> io.BytesIO:
    if len(messages) > _MAX_MESSAGES:
        raise ValueError(f"Export limited to {_MAX_MESSAGES} messages")

    try:
        from weasyprint import HTML
    except (ImportError, OSError):
        raise ExportError("PDF export is not available: weasyprint not installed or system libraries missing")

    safe_title = html_escape(_sanitize_xml_text(title))
    safe_doc = html_escape(_sanitize_xml_text(doc_name))
    date_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    prepared, citations = _prepare_export(messages)
    msg_html = []
    for role, text in prepared:
        safe_content = html_escape(_sanitize_xml_text(text))
        if role == "user":
            msg_html.append(f'<div class="q"><strong>Q:</strong> {safe_content}</div>')
        else:
            msg_html.append(f'<div class="a">{safe_content}</div>')

    refs_html = ""
    if citations:
        refs_items = []
        for c in citations:
            idx = c.get("ref_index", 0)
            refs_items.append(f"<li>[{idx}] {html_escape(_sanitize_xml_text(_format_footnote(c)))}</li>")
        refs_html = f'<h2>References</h2><ul class="references">{"".join(refs_items)}</ul>'

    html_str = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
  @page {{ size: A4; margin: 2cm; }}
  body {{ font-family: 'Noto Sans CJK SC', 'Noto Sans', sans-serif; font-size: 11pt; color: #333; margin: 0; }}
  h1 {{ font-size: 18pt; margin-bottom: 4pt; }}
  .meta {{ color: #888; font-size: 9pt; margin-bottom: 16pt; }}
  .q {{ background: #f5f5f5; padding: 8pt; margin: 6pt 0; border-radius: 4pt; }}
  .a {{ padding: 8pt; margin: 6pt 0; }}
  .q, .a {{ white-space: pre-wrap; overflow-wrap: anywhere; }}
  h2 {{ font-size: 13pt; margin-top: 20pt; break-after: avoid; }}
  .references {{ list-style: none; padding-left: 0; font-size: 9pt; color: #666; }}
  .references li {{ break-inside: avoid; }}
  .footer {{ text-align: center; color: #aaa; font-size: 8pt; margin-top: 30pt; }}
</style></head><body>
<h1>{safe_title}</h1>
<div class="meta">Document: {safe_doc} | Exported: {date_str}</div>
{"".join(msg_html)}
{refs_html}
<div class="footer">Generated by DocTalk &mdash; www.doctalk.site</div>
</body></html>"""

    try:
        buf = io.BytesIO()
        HTML(string=html_str).write_pdf(buf)
        buf.seek(0)
        return buf
    except Exception as e:
        raise ExportError(f"PDF rendering failed: {e}") from e
